import hashlib
import json
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from doctask.domain import SourceSpan


class UnsupportedSourceError(ValueError):
    pass


def _span(
    *,
    source_id: str,
    source_version_id: str,
    text: str,
    locator: dict[str, object],
) -> SourceSpan:
    digest = hashlib.sha256(text.encode()).hexdigest()
    address = json.dumps(locator, sort_keys=True, separators=(",", ":"))
    identity = hashlib.sha256(
        f"{source_version_id}:{address}:{digest}".encode()
    ).hexdigest()
    return SourceSpan(
        span_id=f"span:{identity[:24]}",
        source_id=source_id,
        source_version_id=source_version_id,
        text=text,
        locator=locator,
        content_hash=digest,
    )


def parse_docx(*, source_id: str, source_version_id: str, content: bytes) -> list[SourceSpan]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2_000 or sum(entry.file_size for entry in entries) > 50_000_000:
                raise UnsupportedSourceError("DOCX decompressed content exceeds safety limits")
    except BadZipFile as exc:
        raise UnsupportedSourceError("Source is not a readable DOCX document") from exc
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise UnsupportedSourceError("Source is not a readable DOCX document") from exc

    spans: list[SourceSpan] = []
    if len(document.paragraphs) > 10_000:
        raise UnsupportedSourceError("DOCX paragraph count exceeds safety limits")
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        spans.append(
            _span(
                source_id=source_id,
                source_version_id=source_version_id,
                text=text,
                locator={"kind": "DOCX_PARAGRAPH", "paragraphIndex": paragraph_index},
            )
        )
    return spans


def parse_source(
    *,
    source_id: str,
    source_version_id: str,
    filename: str,
    content: bytes,
) -> list[SourceSpan]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "docx":
        return parse_docx(
            source_id=source_id,
            source_version_id=source_version_id,
            content=content,
        )
    if suffix == "pdf":
        return _parse_pdf(source_id, source_version_id, content)
    if suffix == "md":
        return _parse_markdown(source_id, source_version_id, content)
    if suffix == "txt":
        return _parse_text(source_id, source_version_id, content)
    raise UnsupportedSourceError("Supported source types are PDF, DOCX, Markdown, and TXT")


def _parse_pdf(source_id: str, source_version_id: str, content: bytes) -> list[SourceSpan]:
    try:
        reader = PdfReader(BytesIO(content))
    except Exception as exc:
        raise UnsupportedSourceError("Source is not a readable PDF document") from exc
    if len(reader.pages) > 250:
        raise UnsupportedSourceError("PDF page count exceeds safety limits")
    spans: list[SourceSpan] = []
    extracted_characters = 0
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            lines = (page.extract_text() or "").splitlines()
        except Exception as exc:
            raise UnsupportedSourceError(f"Unable to extract PDF page {page_number}") from exc
        for line in lines:
            text = line.strip()
            if text:
                extracted_characters += len(text)
                if extracted_characters > 5_000_000 or len(spans) >= 10_000:
                    raise UnsupportedSourceError("PDF extracted content exceeds safety limits")
                spans.append(
                    _span(
                        source_id=source_id,
                        source_version_id=source_version_id,
                        text=text,
                        locator={
                            "kind": "PDF_TEXT",
                            "pageNumber": page_number,
                            "excerpt": text,
                        },
                    )
                )
    return spans


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedSourceError("Text source must be UTF-8") from exc


def _parse_markdown(
    source_id: str, source_version_id: str, content: bytes
) -> list[SourceSpan]:
    heading_path: list[str] = []
    spans: list[SourceSpan] = []
    for line_number, raw in enumerate(_decode_text(content).splitlines(), start=1):
        text = raw.strip()
        if not text:
            continue
        if text.startswith("#"):
            marker, separator, heading = text.partition(" ")
            if separator and set(marker) == {"#"}:
                level = len(marker)
                heading_path = heading_path[: level - 1]
                heading_path.append(heading.strip())
        spans.append(
            _span(
                source_id=source_id,
                source_version_id=source_version_id,
                text=text.lstrip("#").strip() if text.startswith("#") else text,
                locator={
                    "kind": "MARKDOWN_LINE",
                    "headingPath": list(heading_path),
                    "lineStart": line_number,
                    "lineEnd": line_number,
                },
            )
        )
    return spans


def _parse_text(source_id: str, source_version_id: str, content: bytes) -> list[SourceSpan]:
    spans: list[SourceSpan] = []
    for line_number, raw in enumerate(_decode_text(content).splitlines(), start=1):
        text = raw.strip()
        if text:
            spans.append(
                _span(
                    source_id=source_id,
                    source_version_id=source_version_id,
                    text=text,
                    locator={
                        "kind": "TEXT_LINE",
                        "lineStart": line_number,
                        "lineEnd": line_number,
                    },
                )
            )
    return spans
