from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent


def build_project_plan() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.start_type = WD_SECTION.NEW_PAGE

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading = document.styles["Heading 1"]
    heading.font.name = "Calibri"
    heading.font.size = Pt(16)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(12)
    title_run = title.add_run("PROJECT LIGHTHOUSE - DELIVERY PLAN")
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)

    document.add_heading("Milestones", level=1)
    document.add_paragraph(
        "Milestone M3: Production readiness - due 18 September 2026 - status at risk."
    )
    document.add_heading("Governance", level=1)
    document.add_paragraph("The program director approves changes to milestone dates.")
    document.save(ROOT / "corpus" / "project-plan.docx")


def build_weekly_status() -> None:
    path = ROOT / "corpus" / "weekly-status.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(72, height - 90, "Project Lighthouse - Weekly Status")
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(72, height - 140, "Risks")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(72, height - 170, "Risk R7: Vendor delay - status open - severity high.")
    pdf.setFont("Helvetica", 9)
    pdf.drawRightString(width - 72, 48, "Page 1")
    pdf.save()


if __name__ == "__main__":
    (ROOT / "corpus").mkdir(parents=True, exist_ok=True)
    build_project_plan()
    build_weekly_status()
